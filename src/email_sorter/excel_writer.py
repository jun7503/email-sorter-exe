# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook, load_workbook

INDEX_SHEET     = "_Index"
TOPICMAP_SHEET  = "_TopicMap"
SUMMARY_SHEET   = "Summary"

TOPIC_COLUMNS = [
    "Received","From","To","Cc","Subject","SuperTopicID","SuperTopicName","TopicID","SubTopicName",
    "Issue","Milestone","Sender","Body","MessageID","Hash","FilePath","ThreadKey","UpdateGroupKey","Status"
]

class ExcelStore:
    """
    Append-only Excel writer/reader.
    Creates base sheets:
      - _Index      (MessageKey|ThreadKey|UpdateGroupKey|TopicSheet|Row)
      - _TopicMap   (SuperTopicID|SuperTopicName|Domain|TopicID|SubTopicName)
      - Summary     (rebuilt at the end of a batch)
      - topic_<domain> sheets for data rows
    """
    def __init__(self, xlsx_path: str):
        self.xlsx_path = xlsx_path
        self.wb = None
        self._open_or_create()

    # ---- init & persistence ----
    def _open_or_create(self):
        p = Path(self.xlsx_path)
        if p.exists():
            self.wb = load_workbook(p)
        else:
            wb = Workbook()
            # remove default "Sheet"
            wb.remove(wb.active)
            wb.create_sheet(INDEX_SHEET)
            wb.create_sheet(TOPICMAP_SHEET)
            wb.create_sheet(SUMMARY_SHEET)
            wb[INDEX_SHEET].append(["MessageKey","ThreadKey","UpdateGroupKey","TopicSheet","Row"])
            wb[TOPICMAP_SHEET].append(["SuperTopicID","SuperTopicName","Domain","TopicID","SubTopicName"])
            wb.save(p)
            self.wb = wb

        # ensure base sheets exist (idempotent)
        for s in (INDEX_SHEET, TOPICMAP_SHEET, SUMMARY_SHEET):
            if s not in self.wb.sheetnames:
                self.wb.create_sheet(s)
        self.save()

    def set_path(self, xlsx_path: str):
        """If user chooses a new Save As path during session."""
        self.xlsx_path = xlsx_path
        self.wb = None
        self._open_or_create()

    def save(self):
        self.wb.save(self.xlsx_path)

    # ---- topic sheets & appends ----
    def topic_sheet(self, domain: str) -> str:
        safe = (domain or "unknown.local").replace("/", "_").replace("\\", "_").replace(":", "_")
        name = f"topic_{safe}"
        if name not in self.wb.sheetnames:
            ws = self.wb.create_sheet(name)
            ws.append(TOPIC_COLUMNS)
            self.save()
        return name

    def append_topic_row(self, domain: str, row: dict):
        ws = self.wb[self.topic_sheet(domain)]
        ws.append([row.get(col, "") for col in TOPIC_COLUMNS])
        self.save()

    def last_row_index(self, domain: str) -> int:
        ws = self.wb[self.topic_sheet(domain)]
        return ws.max_row  # 1-based (includes header row)

    # ---- index operations ----
    def append_index(self, message_key: str, thread_key: str, update_key: str, topic_sheet: str, row_index: int):
        ws = self.wb[INDEX_SHEET]
        ws.append([message_key, thread_key, update_key, topic_sheet, row_index])
        self.save()

    def load_index_keys(self):
        """Return a set of MessageKey values already seen (for idempotency)."""
        ws = self.wb[INDEX_SHEET]
        keys = set()
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1:
                continue
            if not row:
                continue
            if row[0]:
                keys.add(str(row[0]))
        return keys

    def latest_body_for_update_group(self, update_key: str) -> str:
        """Scan bottom-up across topic_* sheets to find the most recent Body for update_key."""
        for name in reversed(self.wb.sheetnames):
            if not name.startswith("topic_"):
                continue
            ws = self.wb[name]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            try:
                i_key = header.index("UpdateGroupKey") + 1
                i_body = header.index("Body") + 1
            except ValueError:
                continue

            for r in range(ws.max_row, 1, -1):
                k = ws.cell(row=r, column=i_key).value
                if k == update_key:
                    return ws.cell(row=r, column=i_body).value or ""
        return ""

    def rebuild_index_from_topics(self):
        """Maintenance: rebuild _Index if user bulk-moved rows or merged files."""
        ws_idx = self.wb[INDEX_SHEET]
        ws_idx.delete_rows(1, ws_idx.max_row)
        ws_idx.append(["MessageKey","ThreadKey","UpdateGroupKey","TopicSheet","Row"])

        for name in self.wb.sheetnames:
            if not name.startswith("topic_"):
                continue
            ws = self.wb[name]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            col = {h: i + 1 for i, h in enumerate(header)}
            for r in range(2, ws.max_row + 1):
                msgid = ws.cell(r, col.get("MessageID", 0)).value or ""
                hsh   = ws.cell(r, col.get("Hash", 0)).value or ""
                message_key = f"MID:{msgid}" if msgid else f"HASH:{hsh}"
                thread_key  = ws.cell(r, col.get("ThreadKey", 0)).value or ""
                update_key  = ws.cell(r, col.get("UpdateGroupKey", 0)).value or ""
                ws_idx.append([message_key, thread_key, update_key, name, r])
        self.save()
